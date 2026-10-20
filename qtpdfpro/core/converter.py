"""Bộ Điều phối Chuyển đổi PDF sang DOCX Đa Chế độ (Multi-Engine PDF-to-DOCX Orchestrator).

Hỗ trợ 5 chế độ chuyển đổi:
- `nd30` (Văn bản hành chính): Tái dựng văn bản đúng thể thức Nghị định 30/2020/NĐ-CP — lề trang,
  giãn dòng, cách đoạn đo trên PDF rồi ước lượng về giá trị gốc; phần đầu & khối ký trong bảng
  không viền; đường kẻ giữ đúng toạ độ PDF; số trang là trường PAGE (`nd30_engine.py`).
- `hybrid`: Kết hợp `pdf2docx` tái tạo cấu trúc (đoạn văn, bảng biểu,
  cột, hình ảnh, màu sắc) + Bộ Hậu xử lý Chuẩn xác OpenXML (`openxml_fixer.py`).
- `exact` (Giữ nguyên vị trí 1:1): Đặt từng dòng chữ vào khung toạ độ tuyệt đối `<w:framePr>`
  kết hợp lớp đồ hoạ nền vector/ảnh phía sau (`exact_layout_engine.py`).
- `word_com` (Microsoft Word Reflow): Gọi engine chuyển đổi gốc của MS Word qua COM rồi
  chạy Bộ Hậu xử lý Chuẩn xác OpenXML (`word_com_bridge.py`).
- `auto` (Tự động thông minh): Văn bản hành chính (có Quốc hiệu – Tiêu ngữ) → `nd30`; còn lại
  phân tích bố cục trang PDF để chọn `hybrid` hoặc `exact` tuỳ độ phức tạp đồ hoạ vector.
"""

from __future__ import annotations

import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, List, Optional

import fitz  # PyMuPDF
from pdf2docx import Converter as Pdf2DocxConverter

from qtpdfpro.core.exact_layout_engine import convert_pdf_exact_layout
from qtpdfpro.core.fidelity import FidelityReport, evaluate_conversion_fidelity
from qtpdfpro.core.admin_layout import analyze_admin_document
from qtpdfpro.core.layout_analyzer import analyze_page_layout
from qtpdfpro.core.layout_verify import LayoutVerifyReport, verify_docx_layout
from qtpdfpro.core.nd30_engine import convert_pdf_nd30
from qtpdfpro.core.nd30_standard import DocxLayoutPlan, ND30Report
from qtpdfpro.core.openxml_fixer import OpenXMLFixReport, fix_docx_package
from qtpdfpro.core.pdf_primitives import (
    extract_widget_overlays as _extract_widget_overlays,
    splice_overlapping_stamped_text as _splice_overlapping_stamped_text,
)
from qtpdfpro.core.word_com_bridge import convert_via_ms_word_com, is_ms_word_available

ProgressCallback = Optional[Callable[[int, int, str], None]]


class ConversionError(Exception):
    """Lỗi tường minh bằng tiếng Việt khi chuyển đổi PDF sang DOCX."""


@dataclass
class ConversionOptions:
    """Tuỳ chọn cấu hình cho một lượt chuyển đổi PDF sang DOCX."""

    mode: str = "hybrid"  # "auto" | "nd30" | "hybrid" | "exact" | "word_com"
    pages: Optional[List[int]] = None  # Danh sách trang 0-based; None = tất cả trang
    default_font: str = "Times New Roman"
    normalize_vietnamese: bool = True
    fix_table_padding: bool = True
    prevent_line_overflow: bool = True
    exact_background_dpi: int = 180
    password: Optional[str] = None
    # Chế độ nd30: "faithful" = giữ đúng bố cục PDF (mặc định), "nd30" = ép chuẩn hoá theo NĐ30
    nd30_policy: str = "faithful"
    # Dựng lại DOCX bằng Microsoft Word và so vị trí từng dòng với PDF (chậm hơn, cần MS Word)
    verify_layout: bool = False


@dataclass
class ConversionResult:
    """Kết quả trả về sau khi hoàn tất chuyển đổi PDF sang DOCX."""

    pdf_path: str
    docx_path: str
    mode_used: str
    pages_converted: int
    elapsed_seconds: float
    fix_report: OpenXMLFixReport = field(default_factory=OpenXMLFixReport)
    fidelity: FidelityReport = field(default_factory=FidelityReport)
    nd30_report: Optional[ND30Report] = None
    nd30_plan: Optional[DocxLayoutPlan] = None
    layout_check: Optional[LayoutVerifyReport] = None
    notes: List[str] = field(default_factory=list)


def is_admin_document(pdf_doc: fitz.Document, page_index: int = 0) -> bool:
    """Nhận diện văn bản hành chính: trang đầu có khối Quốc hiệu – Tiêu ngữ 2 cột."""
    try:
        ad = analyze_admin_document(pdf_doc, [page_index])
        return ad.header is not None and not ad.is_scanned
    except Exception:
        return False


def parse_page_range(spec: str, total_pages: int) -> List[int]:
    """Phân tích chuỗi phạm vi trang (1-based, vd `'1-3, 5, 8-10'`) thành danh sách chỉ số 0-based."""
    if total_pages <= 0:
        return []
    cleaned = (spec or "").strip()
    if not cleaned or cleaned.lower() in ("all", "tatca", "tất cả", "*"):
        return list(range(total_pages))

    indices: set[int] = set()
    for part in cleaned.split(","):
        token = part.strip()
        if not token:
            continue
        if "-" in token:
            pieces = token.split("-", 1)
            try:
                start_1 = int(pieces[0].strip())
                end_1 = int(pieces[1].strip())
            except ValueError as exc:
                raise ConversionError(f"Phạm vi trang không hợp lệ: '{token}'") from exc
            if start_1 > end_1:
                start_1, end_1 = end_1, start_1
            for p1 in range(start_1, end_1 + 1):
                if 1 <= p1 <= total_pages:
                    indices.add(p1 - 1)
        else:
            try:
                p1 = int(token)
            except ValueError as exc:
                raise ConversionError(f"Số trang không hợp lệ: '{token}'") from exc
            if 1 <= p1 <= total_pages:
                indices.add(p1 - 1)

    if not indices:
        raise ConversionError(
            f"Không có trang nào trong phạm vi '{spec}' thuộc tài liệu (1..{total_pages})."
        )
    return sorted(indices)


def convert_pdf_to_docx(
    pdf_path: str | Path,
    docx_path: str | Path,
    options: Optional[ConversionOptions] = None,
    progress_cb: ProgressCallback = None,
) -> ConversionResult:
    """Thực hiện chuyển đổi tệp PDF sang DOCX kèm hậu xử lý chuẩn xác và kiểm định chất lượng."""
    t0 = time.perf_counter()
    opts = options or ConversionOptions()
    src = Path(pdf_path).resolve()
    dst = Path(docx_path).resolve()

    if not src.exists():
        raise ConversionError(f"Không tìm thấy tệp PDF đầu vào: {src}")

    dst.parent.mkdir(parents=True, exist_ok=True)

    # 1. Kiểm tra tệp PDF và xác định danh sách trang cần chuyển
    with fitz.open(str(src)) as doc:
        if doc.needs_pass:
            if not opts.password or not doc.authenticate(opts.password):
                raise ConversionError("Tệp PDF có mật khẩu bảo vệ. Vui lòng nhập đúng mật khẩu.")
        total_pages = len(doc)
        if total_pages == 0:
            raise ConversionError("Tệp PDF không có trang nào.")

        target_pages = (
            [p for p in opts.pages if 0 <= p < total_pages]
            if opts.pages is not None
            else list(range(total_pages))
        )
        if not target_pages:
            target_pages = list(range(total_pages))

        # Xác định chế độ thực tế nếu người dùng chọn "auto"
        actual_mode = opts.mode.lower().strip()
        if actual_mode == "auto" and is_admin_document(doc, target_pages[0]):
            actual_mode = "nd30"
        if actual_mode == "auto":
            sample_pages = target_pages[: min(3, len(target_pages))]
            summaries = [analyze_page_layout(doc[p], include_vectors=False) for p in sample_pages]
            if any(s.recommended_mode == "exact" for s in summaries):
                actual_mode = "exact"
            else:
                actual_mode = "hybrid"

    total_target = len(target_pages)
    if progress_cb:
        progress_cb(0, total_target, f"Khởi động engine chuyển đổi ({actual_mode})...")

    # 2. Chạy Engine chuyển đổi chính
    notes: List[str] = []
    nd30_report: Optional[ND30Report] = None
    nd30_plan: Optional[DocxLayoutPlan] = None
    if actual_mode == "nd30":
        try:
            info = convert_pdf_nd30(
                src,
                dst,
                pages=target_pages,
                policy=opts.nd30_policy,
                password=opts.password,
                progress_cb=progress_cb,
            )
            nd30_report, nd30_plan = info.report, info.plan
            notes.extend(info.notes)
        except Exception as exc:
            notes.append(f"Engine NĐ30 không xử lý được tệp này ({exc}); đã chuyển sang chế độ hybrid.")
            actual_mode = "hybrid"

    if actual_mode == "word_com":
        if is_ms_word_available():
            if progress_cb:
                progress_cb(1, total_target, "Đang chuyển đổi qua Microsoft Word Reflow...")
            ok = convert_via_ms_word_com(src, dst, pages=target_pages)
            if not ok:
                actual_mode = "hybrid"
        else:
            actual_mode = "hybrid"

    if actual_mode == "hybrid":
        try:
            _run_hybrid_pdf2docx(
                src=src,
                dst=dst,
                target_pages=target_pages,
                total_doc_pages=total_pages,
                progress_cb=progress_cb,
            )
        except Exception:
            # Nếu pdf2docx gặp cấu trúc vector dị biệt, tự động chuyển sang chế độ khung tuyệt đối 1:1
            actual_mode = "exact"

    if actual_mode == "exact":
        convert_pdf_exact_layout(
            pdf_path=src,
            docx_path=dst,
            pages=target_pages,
            background_dpi=opts.exact_background_dpi,
            default_font=opts.default_font,
            normalize_vietnamese=opts.normalize_vietnamese,
            progress_cb=progress_cb,
        )

    # 3. Chạy Bộ Hậu xử lý Chuẩn xác OpenXML (QT OpenXML Precision Fixer)
    if progress_cb:
        progress_cb(
            total_target,
            total_target,
            "Đang chuẩn hoá phông chữ, tiếng Việt & tinh chỉnh lề OpenXML...",
        )

    fix_report = fix_docx_package(
        docx_path=dst,
        default_font=opts.default_font,
        normalize_vietnamese=opts.normalize_vietnamese,
        fix_table_padding=(opts.fix_table_padding and actual_mode not in ("exact", "nd30")),
        prevent_line_overflow=(opts.prevent_line_overflow and actual_mode not in ("exact", "nd30")),
        # Engine NĐ30 đã đặt lề trang / khoảng cách đầu trang chuẩn → không tinh chỉnh lại
        tune_page_margins=(actual_mode != "nd30"),
    )

    # 4. Đánh giá độ chính xác (Fidelity Inspector)
    if progress_cb:
        progress_cb(total_target, total_target, "Đang kiểm định độ chính xác tài liệu...")

    fidelity = evaluate_conversion_fidelity(src, dst, pages=target_pages)

    layout_check: Optional[LayoutVerifyReport] = None
    if opts.verify_layout:
        if progress_cb:
            progress_cb(total_target, total_target, "Đang kiểm chứng bố cục bằng Microsoft Word...")
        layout_check = verify_docx_layout(src, dst, pages=target_pages)
    elapsed = round(time.perf_counter() - t0, 3)

    return ConversionResult(
        pdf_path=str(src),
        docx_path=str(dst),
        mode_used=actual_mode,
        pages_converted=total_target,
        elapsed_seconds=elapsed,
        fix_report=fix_report,
        fidelity=fidelity,
        nd30_report=nd30_report,
        nd30_plan=nd30_plan,
        layout_check=layout_check,
        notes=notes,
    )


_PDF2DOCX_PATCHED = False
_ACTIVE_PAGE_OVERLAYS: dict[int, list[tuple[bytes, float, float, float, bool]]] = {}
_ACTIVE_RULE_RECTS: dict[int, list[fitz.Rect]] = {}


def _make_horizontal_rule_png(
    width_pt: float, color_rgb: tuple[float, float, float] = (0.0, 0.0, 0.0)
) -> bytes:
    """Tạo ảnh PNG thanh ngang sắc nét cho đường kẻ dưới tiêu ngữ / tiêu đề hành chính."""
    px_h = 3
    target_h_pt = 0.60
    px_w = max(12, int(round(px_h * width_pt / target_h_pt)))
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, px_w, px_h), True)
    r = int(round(max(0.0, min(1.0, color_rgb[0])) * 255))
    g = int(round(max(0.0, min(1.0, color_rgb[1])) * 255))
    b = int(round(max(0.0, min(1.0, color_rgb[2])) * 255))
    pix.clear_with(0)
    # Tô 2 hàng pixel (độ dày ~0.4pt) để đường kẻ thanh mảnh đúng chuẩn NĐ 30/2020/NĐ-CP
    sub = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 1, px_w, px_h), True)
    sub.set_rect(sub.irect, (r, g, b, 255))
    pix.copy(sub, sub.irect)
    return pix.tobytes("png")


def _extract_standalone_horizontal_rules(
    page: fitz.Page,
) -> tuple[list[tuple[bytes, float, float, float, bool]], list[fitz.Rect]]:
    """Phát hiện các đường kẻ ngang độc lập (như đường kẻ dưới Quốc hiệu - Tiêu ngữ,
    dưới tên cơ quan ban hành và dưới trích yếu Quyết định) để bảo toàn 100% trong Word.
    """
    drawings = page.get_drawings() or []
    if not drawings:
        return [], []

    # Thu thập các đoạn thẳng đứng hoặc hình chữ nhật (biên bảng biểu)
    non_horizontal_rects: list[fitz.Rect] = []
    candidates: list[tuple[float, float, float, tuple[float, float, float], fitz.Rect]] = []

    for d in drawings:
        items = d.get("items") or []
        rect = fitz.Rect(d.get("rect", (0, 0, 0, 0)))
        if len(items) == 1 and items[0][0] == "l":
            p0, p1 = items[0][1], items[0][2]
            dx = abs(p1.x - p0.x)
            dy = abs(p1.y - p0.y)
            if dy <= 1.5 and 15.0 <= dx <= page.rect.width * 0.85:
                color = d.get("color") or (0.0, 0.0, 0.0)
                if not isinstance(color, (tuple, list)) or len(color) < 3:
                    color = (0.0, 0.0, 0.0)
                x0 = min(p0.x, p1.x)
                x1 = max(p0.x, p1.x)
                y_mid = 0.5 * (p0.y + p1.y)
                candidates.append((x0, y_mid, x1 - x0, (color[0], color[1], color[2]), rect))
                continue
        non_horizontal_rects.append(rect)

    overlays: list[tuple[bytes, float, float, float, bool]] = []
    rule_rects: list[fitz.Rect] = []

    for x0, y_mid, width_pt, color_rgb, rect in candidates:
        expanded = fitz.Rect(rect.x0 - 6.0, rect.y0 - 6.0, rect.x1 + 6.0, rect.y1 + 6.0)
        if any(expanded.intersects(nh) for nh in non_horizontal_rects):
            continue
        png_bytes = _make_horizontal_rule_png(width_pt, color_rgb)
        overlays.append((png_bytes, width_pt, x0, y_mid - 0.3, False))
        rule_rects.append(rect)

    return overlays, rule_rects


def _add_floating_overlay(
    p,
    png_bytes: bytes,
    width_pt: float,
    pos_x_pt: float,
    pos_y_pt: float,
    behind_doc: bool = False,
) -> None:
    """Gắn ảnh PNG nổi (`<wp:anchor>`) theo toạ độ trang tuyệt đối vào đoạn văn đầu trang."""
    from io import BytesIO
    from docx.shared import Pt
    from pdf2docx.common.docx import _CT_Anchor

    run = p.add_run()
    r_id, image = run.part.get_or_add_image(BytesIO(png_bytes))
    cx, cy = image.scaled_dimensions(Pt(width_pt), None)
    shape_id, filename = run.part.next_id, image.filename
    anchor = _CT_Anchor.new_pic_anchor(
        shape_id, r_id, filename, cx, cy, Pt(pos_x_pt), Pt(pos_y_pt)
    )
    anchor.set("behindDoc", "1" if behind_doc else "0")
    # Tăng relativeHeight để chữ ký/con dấu nổi rõ trên nền trang
    anchor.set("relativeHeight", "25165824" if not behind_doc else "0")
    run._r.add_drawing(anchor)


def _install_pdf2docx_patches() -> None:
    """Cài đặt bộ bản vá chuyên sâu cho `pdf2docx` phục vụ văn bản hành chính Việt Nam:

    1. Ghép văn bản đóng dấu đè lên khoảng trắng (Số văn bản, Ngày ban hành) vào dòng chủ.
    2. Trích xuất chữ ký số / con dấu đỏ (`page.widgets()`) và đường kẻ ngang độc lập
       (dưới Quốc hiệu - Tiêu ngữ, tên cơ quan, trích yếu) làm lớp phủ PNG trong suốt chuẩn toạ độ.
    3. Sửa lỗi tách cột `RawPage.parse_section` & `RawPage._create_section` để:
       - Tiêu đề căn giữa trang (như `QUYẾT ĐỊNH`) không bị hút nhầm vào cột phải.
       - Khối `Nơi nhận:` nhiều dòng bên trái và khối ký tên bên phải luôn giữ nguyên trong 1 phân đoạn 2 cột.
    4. Sửa lỗi `Lines.parse_line_break` để các dòng trong khối căn giữa (`CENTER`), danh sách gạch đầu dòng
       (`- `, `• `) và dòng kết thúc bằng `:` luôn ngắt dòng chính xác.
    """
    global _PDF2DOCX_PATCHED
    if _PDF2DOCX_PATCHED:
        return
    try:
        from pdf2docx.common import constants
        from pdf2docx.common.Collection import Collection
        from pdf2docx.common.share import TextAlignment
        from pdf2docx.layout.Column import Column
        from pdf2docx.layout.Section import Section
        from pdf2docx.layout.Sections import Sections
        from pdf2docx.page.RawPage import RawPage
        from pdf2docx.page.RawPageFitz import RawPageFitz
        from pdf2docx.text.Lines import Lines

        orig_preprocess_text = RawPageFitz._preprocess_text
        orig_preprocess_shapes = RawPageFitz._preprocess_shapes
        orig_make_docx = Sections.make_docx
        orig_parse_line_break = Lines.parse_line_break

        def _patched_preprocess_text(self, **settings):
            blocks = orig_preprocess_text(self, **settings)
            try:
                blocks = _splice_overlapping_stamped_text(blocks)
                if self.page_engine is not None:
                    pno = self.page_engine.number
                    w_overlays = _extract_widget_overlays(self.page_engine)
                    r_overlays, r_rects = _extract_standalone_horizontal_rules(
                        self.page_engine
                    )
                    _ACTIVE_PAGE_OVERLAYS[pno] = r_overlays + w_overlays
                    _ACTIVE_RULE_RECTS[pno] = r_rects
            except Exception:
                pass
            return blocks

        def _patched_preprocess_shapes(self, **settings):
            shapes, images = orig_preprocess_shapes(self, **settings)
            try:
                if self.page_engine is not None:
                    r_rects = _ACTIVE_RULE_RECTS.get(self.page_engine.number, [])
                    if r_rects:
                        filtered_shapes = []
                        for sh in shapes:
                            if "start" in sh and "end" in sh:
                                sx0, sy0 = sh["start"]
                                sx1, sy1 = sh["end"]
                                sb = fitz.Rect(
                                    min(sx0, sx1),
                                    min(sy0, sy1),
                                    max(sx0, sx1),
                                    max(sy0, sy1),
                                )
                            else:
                                sb = fitz.Rect(sh.get("bbox", (0, 0, 0, 0)))
                            if any(
                                abs(sb.x0 - rr.x0) < 3.5
                                and abs(sb.x1 - rr.x1) < 3.5
                                and abs(sb.y0 - rr.y0) < 3.5
                                for rr in r_rects
                            ):
                                continue
                            filtered_shapes.append(sh)
                        shapes = filtered_shapes
            except Exception:
                pass
            return shapes, images

        def _safe_create_section(num_col: int, elements, h_range: tuple, y_ref: float):
            if not elements:
                return None
            X0, X1 = h_range
            if num_col == 1:
                x0, y0, x1, y1 = elements.bbox
                column = Column().update_bbox((X0, y0, X1, y1))
                column.add_elements(elements)
                section = Section(space=0, columns=[column])
                before_space = y0 - y_ref
            else:
                cols = elements.group_by_columns()
                u0, v0, u1, v1 = cols[0].bbox
                m0, n0, m1, n1 = cols[1].bbox
                u = (u1 + m0) / 2.0
                y_top = min(c.bbox[1] for c in cols)
                y_bot = max(c.bbox[3] for c in cols)

                column_1 = Column().update_bbox((X0, v0, u, y_bot))
                column_1.add_elements(elements)

                column_2 = Column().update_bbox((u, y_top, X1, y_bot))
                column_2.add_elements(elements)

                section = Section(space=0, columns=[column_1, column_2])
                before_space = y_top - y_ref

            section.before_space = round(before_space, 1)
            return section

        def _patched_parse_section(self, **settings):
            X0, Y0, X1, _ = self.working_bbox
            page_center = 0.5 * (X0 + X1)

            elements = Collection()
            elements.extend(self.blocks)
            elements.extend(self.shapes)
            if not elements:
                return []

            lines = Collection()
            sections = []

            def close_section(num_col, elems, y_ref_val):
                if sections and sections[-1].num_cols == num_col == 1:
                    column = sections[-1][0]
                    column.union_bbox(elems)
                    column.add_elements(elems)
                else:
                    sec = self._create_section(num_col, elems, (X0, X1), y_ref_val)
                    if sec:
                        sections.append(sec)

            pre_num_col = 1
            y_ref = Y0
            for row in elements.group_by_rows():
                cols = row.group_by_columns()
                current_num_col = len(cols)

                if current_num_col > 2:
                    current_num_col = 1
                elif current_num_col == 2:
                    u0, v0, u1, v1 = cols[0].bbox
                    m0, n0, m1, n1 = cols[1].bbox
                    x0 = (u1 + m0) / 2.0
                    c1, c2 = x0 - X0, X1 - x0
                    w1, w2 = u1 - u0, m1 - m0
                    f = 2.0
                    if not 1 / f <= c1 / c2 <= f or w1 / c1 < 0.25 or w2 / c2 < 0.25:
                        current_num_col = 1

                if pre_num_col == 2 and current_num_col == 1:
                    prev_cols = lines.group_by_columns()
                    rx0, _, rx1, _ = row.bbox
                    row_center = 0.5 * (rx0 + rx1)

                    if len(prev_cols) >= 2:
                        c0_x1 = prev_cols[0].bbox[2]
                        c1_x0 = prev_cols[1].bbox[0]
                        is_centered_heading = (
                            abs(row_center - page_center) <= 35.0
                            and rx0 < c1_x0 - 8.0
                            and rx1 > c0_x1 + 8.0
                        )
                        if not is_centered_heading and (
                            rx1 <= c1_x0 - 4.0 or rx0 >= c0_x1 + 4.0
                        ):
                            current_num_col = 2
                    else:
                        pos = prev_cols[0].bbox[2]
                        if rx1 <= pos or rx0 > pos:
                            current_num_col = 2

                    if current_num_col == 1:
                        x0, y0, x1, y1 = lines.bbox
                        if y1 - y0 < settings["min_section_height"]:
                            pre_num_col = 1

                elif pre_num_col == 2 and current_num_col == 2:
                    combine = Collection(lines)
                    combine.extend(row)
                    if len(combine.group_by_columns(sorted=False)) == 1:
                        current_num_col = 1

                if current_num_col != pre_num_col:
                    close_section(pre_num_col, lines, y_ref)
                    if sections:
                        y_ref = max(c.bbox[3] for c in sections[-1])
                    lines = Collection(row)
                    pre_num_col = current_num_col
                else:
                    lines.extend(row)

            close_section(current_num_col, lines, y_ref)
            return sections

        def _patched_parse_line_break(
            self,
            bbox,
            line_break_width_ratio: float,
            line_break_free_space_ratio: float,
        ):
            orig_parse_line_break(
                self, bbox, line_break_width_ratio, line_break_free_space_ratio
            )
            try:
                block = self.parent
                rows = self.group_by_physical_rows()
                if len(rows) <= 1:
                    return
                for i in range(len(rows) - 1):
                    cur_txt = "".join(ln.text for ln in rows[i]).strip()
                    nxt_txt = "".join(ln.text for ln in rows[i + 1]).lstrip()
                    is_bullet_line = cur_txt.startswith(("-", "–", "•", "+"))
                    if (
                        block.alignment == TextAlignment.CENTER
                        or nxt_txt.startswith(("-", "–", "•", "+"))
                        or (is_bullet_line and cur_txt.endswith(";"))
                        or cur_txt.endswith((":", "./."))
                    ):
                        rows[i][-1].line_break = 1
            except Exception:
                pass

        def _patched_sections_make_docx(self, doc):
            if not self:
                return
            n = len(doc.paragraphs)
            orig_make_docx(self, doc)
            try:
                page_id = getattr(self.parent, "id", -1)
                overlays = _ACTIVE_PAGE_OVERLAYS.get(page_id, [])
                if overlays and n < len(doc.paragraphs):
                    anchor_para = doc.paragraphs[n]
                    for png_bytes, width_pt, pos_x_pt, pos_y_pt, behind_doc in overlays:
                        _add_floating_overlay(
                            anchor_para,
                            png_bytes=png_bytes,
                            width_pt=width_pt,
                            pos_x_pt=pos_x_pt,
                            pos_y_pt=pos_y_pt,
                            behind_doc=behind_doc,
                        )
            except Exception:
                pass

        RawPageFitz._preprocess_text = _patched_preprocess_text
        RawPageFitz._preprocess_shapes = _patched_preprocess_shapes
        RawPage._create_section = staticmethod(_safe_create_section)
        RawPage.parse_section = _patched_parse_section
        Lines.parse_line_break = _patched_parse_line_break
        Sections.make_docx = _patched_sections_make_docx
        _PDF2DOCX_PATCHED = True
    except Exception:
        pass


def _run_hybrid_pdf2docx(
    src: Path,
    dst: Path,
    target_pages: List[int],
    total_doc_pages: int,
    progress_cb: ProgressCallback = None,
) -> None:
    """Chạy `pdf2docx` để phân tích cấu trúc đoạn văn, cột và bảng biểu."""
    import contextlib
    import io as _io
    import logging

    _install_pdf2docx_patches()
    _ACTIVE_PAGE_OVERLAYS.clear()
    _ACTIVE_RULE_RECTS.clear()

    total_target = len(target_pages)
    if progress_cb:
        progress_cb(
            0,
            total_target,
            f"Đang phân tích cấu trúc đoạn văn & bảng biểu ({total_target} trang)...",
        )

    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    with contextlib.redirect_stdout(_io.StringIO()), contextlib.redirect_stderr(_io.StringIO()):
        cv = Pdf2DocxConverter(str(src))
        try:
            is_all = len(target_pages) == total_doc_pages
            if is_all:
                cv.convert(str(dst))
            else:
                cv.convert(str(dst), pages=target_pages)
        finally:
            cv.close()
            _ACTIVE_PAGE_OVERLAYS.clear()
            _ACTIVE_RULE_RECTS.clear()

    if progress_cb:
        progress_cb(total_target, total_target, "Đã dựng xong cấu trúc DOCX cơ bản.")
