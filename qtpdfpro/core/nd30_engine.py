"""Engine chuyển PDF → DOCX chuẩn thể thức Nghị định 30/2020/NĐ-CP (chế độ `nd30`).

Thay vì sao chép hình học từng khối như `pdf2docx` (dựng phần đầu bằng nhiều section chia cột,
số trang là chữ gõ tay, lề trang tính theo mép chữ), engine này **tái dựng văn bản như người soạn
thảo đã làm trong Word**:

- Lề trang, cỡ chữ, khoảng cách dòng, khoảng cách đoạn, thụt đầu dòng được **đo trên PDF rồi ước
  lượng về giá trị gốc** (vd 29,98 mm → 30 mm; 24,13 pt / 14 pt → 1,5 lines) để DOCX dàn trang
  trùng khớp PDF; chính sách `nd30` (tuỳ chọn) kẹp thêm vào khoảng NĐ30 cho phép.
- Phần đầu (Tên cơ quan | Quốc hiệu – Tiêu ngữ; Số, ký hiệu | Địa danh, ngày) và khối ký
  (Nơi nhận | Quyền hạn, chức vụ, họ tên) đặt trong **bảng 2 cột không viền**; từng dòng được
  căn tâm đúng toạ độ PDF bằng thụt lề đoạn trong ô.
- Đường kẻ ngang giữ **đúng toạ độ x và độ dài như PDF**, dựng bằng viền dưới của đoạn văn
  (chỉnh sửa được trong Word, không dùng ảnh).
- Số trang là trường `PAGE` ở đầu trang, không hiển thị trang 1 (đúng NĐ30).
- Con dấu đỏ, chữ ký số giữ nguyên ảnh PNG trong suốt, neo theo đoạn văn gần nhất.
"""

from __future__ import annotations

import statistics
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Callable, List, Optional, Sequence

import fitz  # PyMuPDF
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor

from qtpdfpro.core.admin_layout import (
    PT_PER_MM,
    AdminDocument,
    ImageBlock,
    ImageItem,
    ParaBlock,
    RuleBlock,
    TableBlock,
    TwoColBlock,
    _build_flow_paragraphs,
    analyze_admin_document,
    font_line_metrics,
)
from qtpdfpro.core.nd30_standard import (
    DocxLayoutPlan,
    ND30Report,
    build_docx_plan,
    check_nd30_compliance,
    normalize_policy,
    resolve_line_spacing,
)

ProgressCallback = Optional[Callable[[int, int, str], None]]

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_ND30_CONVERTIBLE_FONTS = {"Times New Roman", "Arial", "Calibri", "Cambria", "Tahoma", "Verdana", "Segoe UI"}

# Thứ tự phần tử con của <w:pPr> theo lược đồ OOXML (để chèn pBdr đúng vị trí)
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_PPR_AFTER_SNAP = _PPR_AFTER_PBDR[_PPR_AFTER_PBDR.index("w:spacing"):]


@dataclass
class ND30ConversionInfo:
    """Thông tin trả về sau khi dựng DOCX bằng engine NĐ30."""

    report: ND30Report
    plan: DocxLayoutPlan
    is_admin: bool
    blocks: int
    header_table: bool
    sign_table: bool
    rules: int
    images: int
    notes: List[str] = field(default_factory=list)


def analyze_pdf_nd30(
    pdf_path: str | Path,
    pages: Optional[Sequence[int]] = None,
    password: Optional[str] = None,
) -> tuple[AdminDocument, ND30Report]:
    """Phân tích & đối chiếu thể thức NĐ30 của một tệp PDF (không tạo DOCX)."""
    with fitz.open(str(pdf_path)) as doc:
        if doc.needs_pass and (not password or not doc.authenticate(password)):
            raise PermissionError("Tệp PDF có mật khẩu bảo vệ.")
        ad = analyze_admin_document(doc, pages)
    return ad, check_nd30_compliance(ad)


def convert_pdf_nd30(
    pdf_path: str | Path,
    docx_path: str | Path,
    pages: Optional[Sequence[int]] = None,
    policy: str = "faithful",
    password: Optional[str] = None,
    progress_cb: ProgressCallback = None,
) -> ND30ConversionInfo:
    """Chuyển PDF văn bản hành chính sang DOCX chuẩn thể thức NĐ30 nhưng giữ đúng bố cục PDF."""
    policy = normalize_policy(policy)
    total = len(pages) if pages is not None else 0
    if progress_cb:
        progress_cb(0, max(1, total), "Đang phân tích bố cục văn bản hành chính (NĐ30)...")

    with fitz.open(str(pdf_path)) as doc:
        if doc.needs_pass and (not password or not doc.authenticate(password)):
            raise PermissionError("Tệp PDF có mật khẩu bảo vệ.")
        ad = analyze_admin_document(doc, pages)

    if ad.is_scanned:
        raise ValueError("Tệp PDF dạng ảnh quét — engine NĐ30 cần lớp chữ (hãy dùng OCR trước).")

    report = check_nd30_compliance(ad)
    plan = build_docx_plan(ad, policy)
    if progress_cb:
        progress_cb(max(1, total) // 2, max(1, total), "Đang dựng DOCX: bảng không viền, đường kẻ, lề chuẩn...")

    writer = _DocxWriter(ad, plan)
    writer.build()
    out = Path(docx_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    writer.document.save(str(out))

    if progress_cb:
        progress_cb(max(1, total), max(1, total), "Đã dựng xong DOCX chuẩn NĐ30.")

    return ND30ConversionInfo(
        report=report,
        plan=plan,
        is_admin=ad.is_admin,
        blocks=len(ad.blocks),
        header_table=ad.header is not None,
        sign_table=ad.sign is not None,
        rules=writer.rule_count,
        images=writer.image_count,
        notes=list(plan.notes),
    )


# ─── Bộ ghi DOCX ───────────────────────────────────────────────────────────


@dataclass
class _Anchor:
    page: int
    top: float
    paragraph: object


class _DocxWriter:
    def __init__(self, ad: AdminDocument, plan: DocxLayoutPlan) -> None:
        self.ad = ad
        self.plan = plan
        self.document = Document()
        self.rule_count = 0
        self.image_count = 0
        self._anchors: list[_Anchor] = []
        self._docpr_id = 100
        self._dominant_spacing = plan.para_spacing_src_pt if plan.para_spacing_src_pt is not None else plan.para_spacing_pt
        self._first_indent_src = (plan.first_indent_src_mm or plan.first_indent_mm) * PT_PER_MM
        self._page_size = (ad.pages[0].width, ad.pages[0].height) if ad.pages else (595.3, 841.9)

    # ── Khung tài liệu ──
    @property
    def text_left(self) -> float:
        return self.plan.margin_left_mm * PT_PER_MM

    @property
    def text_right(self) -> float:
        return self._page_size[0] - self.plan.margin_right_mm * PT_PER_MM

    def _setup_document(self) -> None:
        d = self.document
        styles = d.styles
        normal = styles["Normal"]
        normal.font.name = self.plan.body_font
        normal.font.size = Pt(self.plan.body_size)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), self.plan.body_font)
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "vi-VN")
        pf = normal.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.widow_control = True

        # docDefaults: phông & cỡ chữ cho mọi phần tử
        st_el = styles.element
        doc_defaults = st_el.find(qn("w:docDefaults"))
        if doc_defaults is not None:
            for rf in doc_defaults.iter(qn("w:rFonts")):
                for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    rf.attrib.pop(qn(attr), None)
                for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                    rf.set(qn(attr), self.plan.body_font)
            for sp in doc_defaults.iter(qn("w:spacing")):
                sp.set(qn("w:after"), "0")
                sp.set(qn("w:line"), "240")
                sp.set(qn("w:lineRule"), "auto")

        self._set_modern_compat()
        sec = d.sections[0]
        self._apply_section_geometry(sec, self._page_size)
        if self.plan.add_page_numbers:
            self._add_page_number_header(sec)

    def _set_modern_compat(self) -> None:
        """Bật chế độ bố cục Word 2013+ (compatibilityMode 15) như văn bản soạn bằng Word hiện đại.

        Mẫu mặc định của python-docx dùng chế độ Word 2010 (14) với thuật toán ngắt dòng / căn đều
        khác, làm vài dòng bị đẩy chữ cuối xuống dòng dưới so với PDF gốc.
        """
        settings = self.document.settings.element
        compat = settings.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat")
            settings.append(compat)
        for child in list(compat):
            if child.tag == qn("w:useFELayout"):
                compat.remove(child)
        found = False
        for cs in compat.findall(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                cs.set(qn("w:val"), "15")
                found = True
        if not found:
            compat.append(parse_xml(
                f'<w:compatSetting {nsdecls("w")} w:name="compatibilityMode" '
                'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>'
            ))

    def _apply_section_geometry(self, sec, size: tuple[float, float]) -> None:
        w, h = size
        landscape = w > h
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width = Pt(w)
        sec.page_height = Pt(h)
        top, bottom = self.plan.margin_top_mm, self.plan.margin_bottom_mm
        left, right = self.plan.margin_left_mm, self.plan.margin_right_mm
        if landscape:
            # Trang ngang (phụ lục, bảng lớn): lề trái/phải đổi vai với trên/dưới
            top, bottom, left, right = left, right, bottom, top
        sec.top_margin = Pt(top * PT_PER_MM)
        sec.bottom_margin = Pt(bottom * PT_PER_MM)
        sec.left_margin = Pt(left * PT_PER_MM)
        sec.right_margin = Pt(right * PT_PER_MM)
        sec.header_distance = Pt(self.plan.header_distance_mm * PT_PER_MM)
        sec.footer_distance = Pt(10.0 * PT_PER_MM)
        sec.gutter = Pt(0)

    def _add_page_number_header(self, sec) -> None:
        pages_with_number = set(self.ad.measured.page_number_pages)
        first_idx = self.ad.pages[0].index if self.ad.pages else 0
        hide_first = (first_idx not in pages_with_number) or self.plan.policy == "nd30"
        sec.different_first_page_header_footer = hide_first
        hdr = sec.header
        hdr.is_linked_to_previous = False
        p = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        size = self.plan.body_size

        def fld(kind: str) -> None:
            r = p.add_run()
            self._style_run(r, self.plan.body_font, size, False, False, 0)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)

        fld("begin")
        r = p.add_run()
        self._style_run(r, self.plan.body_font, size, False, False, 0)
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        r._r.append(instr)
        fld("separate")
        r = p.add_run("2")
        self._style_run(r, self.plan.body_font, size, False, False, 0)
        fld("end")
        if hide_first:
            fp = sec.first_page_header
            fp.is_linked_to_previous = False

    # ── Xây dựng ──
    def build(self) -> None:
        self._setup_document()
        body = self.document
        prev_page: Optional[int] = None
        prev_page_bottom = 0.0
        cur_size = self._page_size
        for idx, blk in enumerate(self.ad.blocks):
            page = blk.page
            page_model = self._page(page)
            if page_model is not None and (page_model.width, page_model.height) != cur_size and prev_page is not None:
                cur_size = (page_model.width, page_model.height)
                new_sec = self.document.add_section(WD_SECTION.NEW_PAGE)
                self._apply_section_geometry(new_sec, cur_size)
            new_page = prev_page is not None and page != prev_page
            force_break = new_page and self._is_manual_page_break(prev_page, prev_page_bottom)

            if isinstance(blk, ParaBlock):
                p = self._write_para(body, blk, top_level=True, at_page_top=new_page)
                if force_break:
                    p.paragraph_format.page_break_before = True
            elif isinstance(blk, RuleBlock):
                self._write_rule(body, blk, self.text_left, self.text_right)
            elif isinstance(blk, TwoColBlock):
                self._write_two_col(blk)
            elif isinstance(blk, TableBlock):
                self._carry_gap_before_table(blk.before)
                self._write_data_table(blk)
            elif isinstance(blk, ImageBlock):
                self._write_inline_image(blk)
            prev_page = getattr(blk, "end_page", page)
            prev_page_bottom = getattr(blk, "bottom", prev_page_bottom)
        self._place_floating_images()
        self._remove_leading_empty_paragraph()

    def _page(self, index: int):
        for m in self.ad.pages:
            if m.index == index:
                return m
        return None

    def _is_manual_page_break(self, prev_page: int, prev_bottom: float) -> bool:
        m = self._page(prev_page)
        if m is None:
            return False
        usable_bottom = m.height - self.plan.margin_bottom_mm * PT_PER_MM
        usable_top = self.plan.margin_top_mm * PT_PER_MM
        filled = (prev_bottom - usable_top) / max(1.0, usable_bottom - usable_top)
        return filled < 0.6

    def _remove_leading_empty_paragraph(self) -> None:
        body = self.document.element.body
        first = body[0] if len(body) else None
        if first is not None and first.tag == qn("w:p") and not first.xpath(".//w:t") and not first.xpath(".//w:drawing"):
            if first.find(qn("w:pPr")) is None or first.find(qn("w:pPr")).find(qn("w:sectPr")) is None:
                body.remove(first)

    # ── Định dạng chung ──
    def _font_for(self, family: str) -> str:
        if self.plan.policy == "nd30" and family in _ND30_CONVERTIBLE_FONTS:
            return "Times New Roman"
        return family

    def _size_for(self, size: float, role: str) -> float:
        s = round(size * 2) / 2
        if self.plan.policy == "nd30" and role == "body" and abs(s - self.plan.body_size_src) < 0.3:
            return float(self.plan.body_size)
        return s

    @staticmethod
    def _style_run(run, font: str, size: float, bold: bool, italic: bool, color: int) -> None:
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = bool(bold)
        run.italic = bool(italic)
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font)
        if color and _not_black(color):
            run.font.color.rgb = RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)

    def _set_snap_off(self, p) -> None:
        ppr = p._p.get_or_add_pPr()
        if ppr.find(qn("w:snapToGrid")) is None:
            el = OxmlElement("w:snapToGrid")
            el.set(qn("w:val"), "0")
            ppr.insert_element_before(el, *_PPR_AFTER_SNAP)

    def _line_spacing(self, p, para: ParaBlock, role: str) -> float:
        """Đặt giãn dòng cho đoạn; trả về chiều cao dòng (pt) mà Word sẽ dùng."""
        pf = p.paragraph_format
        rule, value = para.line_rule, para.line_value
        if self.plan.policy == "nd30":
            size = self._size_for(para.size, role)
            lh, _desc = font_line_metrics(self._font_for(para.font))
            pitch = para.line_height() * size / max(0.1, para.size)
            rule, value = resolve_line_spacing(pitch, size, lh, "nd30", is_body=(role == "body"))
        if rule == "auto":
            pf.line_spacing = value
            return value
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(value)
        return value

    def _spacing_before(self, before: float, role: str, at_page_top: bool) -> float:
        d = self._dominant_spacing
        if role == "body":
            if at_page_top:
                before = d if d is not None else before
            elif d is not None and abs(before - d) <= 0.7:
                before = d
            if self.plan.policy == "nd30" and d is not None and abs(before - d) <= 0.7:
                before = max(before, self.plan.para_spacing_pt)
        if before < 0.6:
            return 0.0
        return round(before * 2) / 2

    # ── Đoạn văn ──
    def _write_para(self, container, para: ParaBlock, *, top_level: bool, at_page_top: bool = False,
                    region: Optional[tuple[float, float]] = None, use_first: bool = False):
        p = container.paragraphs[0] if use_first else container.add_paragraph()
        role = para.role
        pf = p.paragraph_format
        pf.alignment = _ALIGN.get(para.align, WD_ALIGN_PARAGRAPH.LEFT)
        pf.space_after = Pt(0)
        pf.space_before = Pt(self._spacing_before(para.before, role, at_page_top))
        self._line_spacing(p, para, role)

        if region is None:
            ind_left, ind_right, first = para.ind_left, para.ind_right, para.first_line
            if role == "body" and abs(first - self._first_indent_src) <= 1.5 * PT_PER_MM:
                first = self.plan.first_indent_mm * PT_PER_MM
            if para.align == "center":
                ind_left = ind_right = 0.0
        else:
            ind_left, ind_right, first = self._cell_indents(para, region)
        pf.left_indent = Pt(round(ind_left, 1)) if abs(ind_left) >= 0.3 else Pt(0)
        pf.right_indent = Pt(round(ind_right, 1)) if abs(ind_right) >= 0.3 else Pt(0)
        pf.first_line_indent = Pt(round(first, 1)) if abs(first) >= 0.3 else Pt(0)
        self._set_snap_off(p)

        for li, line in enumerate(para.lines):
            if li > 0:
                if para.joiner == "break":
                    p.add_run().add_break()
                else:
                    last = p.runs[-1] if p.runs else None
                    if last is not None and not last.text.endswith(" ") and not line.text.startswith(" "):
                        last.text += " "
            for r in line.runs:
                run = p.add_run(r.text)
                self._style_run(run, self._font_for(r.font), self._size_for(r.size, role), r.bold, r.italic, r.color)
        if top_level:
            self._anchors.append(_Anchor(para.page, para.top, p))
        return p

    def _cell_indents(self, para: ParaBlock, region: tuple[float, float]) -> tuple[float, float, float]:
        """Thụt lề đoạn trong ô sao cho tâm / mép dòng trùng toạ độ PDF."""
        a, b = region
        widths = [l.width for l in para.lines]
        if para.align == "center":
            c = statistics.median(l.center for l in para.lines)
            half = min(c - a, b - c)
            need = 0.5 * max(widths) + 3.0
            if half >= need:
                return max(0.0, c - half - a), max(0.0, b - (c + half)), 0.0
            # Không đủ chỗ để căn tâm chính xác → căn trái đúng mép chữ PDF (tránh Word tự ngắt dòng)
            para.align = "left"
            x0 = min(l.x0 for l in para.lines)
            return max(0.0, min(x0 - a, (b - a) - max(widths) - 3.0)), 0.0, 0.0
        if para.align == "right":
            x1 = max(l.x1 for l in para.lines)
            return 0.0, max(0.0, b - x1), 0.0
        x0_first = para.lines[0].x0
        x0_rest = statistics.median(l.x0 for l in para.lines[1:]) if len(para.lines) > 1 else x0_first
        return max(0.0, x0_rest - a), 0.0, x0_first - x0_rest

    # ── Đường kẻ (viền dưới của đoạn rỗng) ──
    def _write_rule(self, container, blk: RuleBlock, a: float, b: float, use_first: bool = False):
        p = container.paragraphs[0] if use_first else container.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(max(1.0, round(blk.height * 4) / 4))
        pf.left_indent = Pt(round(max(0.0, blk.rule.x0 - a), 1))
        pf.right_indent = Pt(round(max(0.0, b - blk.rule.x1), 1))
        ppr = p._p.get_or_add_pPr()
        sz = int(min(24, max(2, round(blk.rule.thickness * 8))))
        r, g, bl = (int(round(max(0.0, min(1.0, c)) * 255)) for c in blk.rule.color)
        pbdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="0" '
            f'w:color="{r:02X}{g:02X}{bl:02X}"/></w:pBdr>'
        )
        ppr.insert_element_before(pbdr, *_PPR_AFTER_PBDR)
        self._set_snap_off(p)
        # Dấu đoạn cỡ 1pt để Word không nới chiều cao dòng
        mark_rpr = ppr.find(qn("w:rPr"))
        if mark_rpr is None:
            mark_rpr = OxmlElement("w:rPr")
            ppr.append(mark_rpr)
        for tag, val in (("w:sz", "2"), ("w:szCs", "2")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            mark_rpr.append(el)
        self.rule_count += 1
        return p

    # ── Bảng 2 cột không viền ──
    def _write_two_col(self, blk: TwoColBlock) -> None:
        d = self.document
        floating_here = [
            im for im in self.ad.floating_images if im.page == blk.page and blk.top - 2 <= im.bbox[1] <= blk.bottom + 40
        ]
        has_anchor = any(a.page == blk.page for a in self._anchors)
        before = blk.before
        if floating_here and not has_anchor:
            # Đoạn neo 1pt trước bảng để gắn con dấu / chữ ký theo đúng trang
            holder = d.add_paragraph()
            hpf = holder.paragraph_format
            hpf.space_before = Pt(round(before * 2) / 2 if before >= 0.6 else 0)
            hpf.space_after = Pt(0)
            hpf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            hpf.line_spacing = Pt(1)
            self._set_snap_off(holder)
            self._anchors.append(_Anchor(blk.page, blk.top - 1.0, holder))
            before = 0.0

        a, mid, b = blk.x_start, blk.x_mid, blk.x_end
        widths = [max(20.0, mid - a), max(20.0, b - mid)]
        table = d.add_table(rows=len(blk.rows), cols=2)
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        for child in list(tbl_pr):
            if child.tag in (qn("w:tblStyle"), qn("w:tblW"), qn("w:tblLook")):
                tbl_pr.remove(child)
        total_tw = int(round(sum(widths) * 20))
        ind_tw = int(round((a - self.text_left) * 20))
        tbl_pr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_tw}" w:type="dxa"/>'))
        tbl_pr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{ind_tw}" w:type="dxa"/>'))
        tbl_pr.append(
            parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                + "".join(f'<w:{s} w:val="nil"/>' for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
                + "</w:tblBorders>"
            )
        )
        tbl_pr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tbl_pr.append(
            parse_xml(
                f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
                f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'
            )
        )
        tbl_pr.append(parse_xml(f'<w:tblLook {nsdecls("w")} w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"/>'))
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(round(w * 20))))

        regions = [(a, mid), (mid, b)]
        for ri, (lcol, rcol) in enumerate(blk.rows):
            row = table.rows[ri]
            for ci, items in enumerate((lcol, rcol)):
                cell = row.cells[ci]
                cell.width = Pt(widths[ci])
                if not items:
                    continue
                for k, it in enumerate(items):
                    item_before = it.before + (before if ri == 0 and k == 0 else 0.0)
                    if isinstance(it, RuleBlock):
                        self._write_rule(cell, it, *regions[ci], use_first=(k == 0))
                    else:
                        saved = it.before
                        it.before = item_before
                        self._write_para(cell, it, top_level=False, region=regions[ci], use_first=(k == 0))
                        it.before = saved

    def _carry_gap_before_table(self, gap: float) -> None:
        """Word không có "khoảng cách trước bảng" → chuyển thành khoảng cách sau của đoạn liền trước."""
        if gap < 0.6:
            return
        body = self.document.element.body
        prev = body[-2] if len(body) >= 2 and body[-1].tag == qn("w:sectPr") else (body[-1] if len(body) else None)
        if prev is None or prev.tag != qn("w:p"):
            return
        from docx.text.paragraph import Paragraph as _P

        para = _P(prev, self.document._body)
        para.paragraph_format.space_after = Pt(round(gap * 2) / 2)

    # ── Bảng dữ liệu có viền ──
    def _write_data_table(self, blk: TableBlock) -> None:
        t = blk.table
        edges = t.col_edges
        if len(edges) < 2 or not t.rows:
            return
        n_cols = len(edges) - 1
        n_rows = len(t.rows)
        d = self.document

        # Lề trong ô đo từ PDF: khoảng cách từ mép ô tới chữ của các ô căn trái
        pads: list[float] = []
        for row in t.rows:
            for c in row:
                if c is not None and c.lines:
                    gap = min(l.x0 for l in c.lines) - c.bbox[0]
                    if 0.0 <= gap <= 14.0:
                        pads.append(gap)
        pad = min(pads) if pads else 5.4
        pad = max(0.0, min(pad, 10.0))

        # Khoảng cách trước đoạn văn trước bảng được đặt vào đoạn rỗng không có — dùng spacing của đoạn sau
        table = d.add_table(rows=n_rows, cols=n_cols)
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        for child in list(tbl_pr):
            if child.tag in (qn("w:tblStyle"), qn("w:tblW")):
                tbl_pr.remove(child)
        widths = [edges[i + 1] - edges[i] for i in range(n_cols)]
        pad_tw = int(round(pad * 20))
        tbl_pr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(round(sum(widths) * 20))}" w:type="dxa"/>'))
        tbl_pr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{int(round((edges[0] - self.text_left) * 20))}" w:type="dxa"/>'))
        tbl_pr.append(
            parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                + "".join(
                    f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                    for s in ("top", "left", "bottom", "right", "insideH", "insideV")
                )
                + "</w:tblBorders>"
            )
        )
        tbl_pr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tbl_pr.append(
            parse_xml(
                f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:left w:w="{pad_tw}" w:type="dxa"/>'
                f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="{pad_tw}" w:type="dxa"/></w:tblCellMar>'
            )
        )
        for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(round(w * 20))))
        for r_i in range(n_rows):
            for c_i in range(n_cols):
                table.cell(r_i, c_i).width = Pt(widths[c_i])

        def col_index(x: float) -> int:
            return min(range(len(edges)), key=lambda i: abs(edges[i] - x))

        row_y = []
        for row in t.rows:
            ys = [c.bbox for c in row if c is not None]
            row_y.append((min(b[1] for b in ys), min(b[3] for b in ys)) if ys else (0.0, 0.0))

        # Chiều cao hàng tối thiểu = khoảng cách 2 đường viền trên PDF trừ độ dày viền (Word tự cộng
        # thêm nét viền vào chiều cao hàng) → giữ nguyên vị trí các dòng sau bảng
        for ri, (y0, y1) in enumerate(row_y):
            if y1 > y0:
                tr_pr = table.rows[ri]._tr.get_or_add_trPr()
                h_tw = int(round((y1 - y0 - 0.5) * 20))
                tr_pr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{h_tw}" w:hRule="atLeast"/>'))

        size_pitch: dict = {}
        for ri, row in enumerate(t.rows):
            for c in row:
                if c is None:
                    continue
                c0 = col_index(c.bbox[0])
                c1 = max(c0 + 1, col_index(c.bbox[2])) - 1
                r1 = ri
                while r1 + 1 < n_rows and row_y[r1 + 1][0] < c.bbox[3] - 2.0:
                    r1 += 1
                try:
                    cell = table.cell(ri, c0)
                    if c1 > c0 or r1 > ri:
                        cell = cell.merge(table.cell(r1, min(c1, n_cols - 1)))
                except Exception:
                    cell = table.cell(ri, c0)
                region = (c.bbox[0] + pad, c.bbox[2] - pad)
                paras = _build_flow_paragraphs(c.lines, region[0], region[1], size_pitch)
                ref = c.bbox[1] + 0.5
                for k, para in enumerate(paras):
                    para.role = "cell"
                    para.resolve_spacing(None)
                    para.before = max(0.0, para.top - ref)
                    ref = para.bottom
                    self._write_para(cell, para, top_level=False, region=region, use_first=(k == 0))

    # ── Ảnh ──
    def _write_inline_image(self, blk: ImageBlock) -> None:
        img = blk.image
        p = self.document.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(self._spacing_before(blk.before, "image", False))
        pf.space_after = Pt(0)
        cx = 0.5 * (img.bbox[0] + img.bbox[2])
        if abs(cx - 0.5 * (self.text_left + self.text_right)) <= 6.0:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            pf.left_indent = Pt(max(0.0, img.bbox[0] - self.text_left))
        run = p.add_run()
        try:
            run.add_picture(BytesIO(img.png), width=Pt(img.bbox[2] - img.bbox[0]))
            self.image_count += 1
        except Exception:
            pass
        self._anchors.append(_Anchor(blk.page, img.bbox[1], p))

    def _place_floating_images(self) -> None:
        for img in self.ad.floating_images:
            cands = [a for a in self._anchors if a.page == img.page]
            if not cands:
                cands = [a for a in self._anchors if a.page <= img.page] or self._anchors
                if not cands:
                    continue
                anchor = cands[-1]
                v_rel, v_off = "page", img.bbox[1]
            else:
                above = [a for a in cands if a.top <= img.bbox[1] + 1.0]
                anchor = above[-1] if above else cands[0]
                v_rel, v_off = "paragraph", img.bbox[1] - anchor.top
            self._add_anchor_image(anchor.paragraph, img, v_rel, v_off)

    def _add_anchor_image(self, paragraph, img: ImageItem, v_rel: str, v_off: float) -> None:
        try:
            run = paragraph.add_run()
            r_id, image = run.part.get_or_add_image(BytesIO(img.png))
        except Exception:
            return
        w_pt = img.bbox[2] - img.bbox[0]
        h_pt = img.bbox[3] - img.bbox[1]
        cx, cy = int(Emu(Pt(w_pt))), int(Emu(Pt(h_pt)))
        x_emu = int(Emu(Pt(img.bbox[0])))
        y_emu = int(Emu(Pt(v_off)))
        self._docpr_id += 1
        did = self._docpr_id
        xml = (
            f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{251660000 + did}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{x_emu}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="{v_rel}"><wp:posOffset>{y_emu}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{cx}" cy="{cy}"/>'
            '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
            "<wp:wrapNone/>"
            f'<wp:docPr id="{did}" name="QTPdfPro Overlay {did}"/>'
            '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{did}" name="{image.filename}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{r_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
            "</a:graphicData></a:graphic></wp:anchor>"
        )
        drawing = OxmlElement("w:drawing")
        drawing.append(parse_xml(xml))
        run._r.append(drawing)
        self.image_count += 1


def _not_black(color: int) -> bool:
    r, g, b = (color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF
    return max(r, g, b) > 0x30

